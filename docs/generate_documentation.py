from pathlib import Path
import json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
out=root/'docs'/'NFC_Card_Lab_Dokumentacja_projektu_2026-10-05.docx'
doc=Document()
sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
sec.header_distance=Cm(.7); sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(7)
normal=doc.styles['Normal']; normal.font.size=Pt(10.5); normal.paragraph_format.line_spacing=1.08
normal.paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(30)
doc.styles['Subtitle'].font.size=Pt(14)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(12)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0]; header.text='NFC CARD LAB   |   Dokumentacja projektu   |   5 października 2026'
header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Wersja dokumentacji 1.1   •   Strona ').font.size=Pt(8)
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
doc.core_properties.title='NFC Card Lab Dokumentacja projektu'
doc.core_properties.subject='Instrukcja obsługi i dokumentacja techniczna Android NFC'
doc.core_properties.author='MisiekLab'
doc.core_properties.keywords='NFC, Android, Kotlin, MIFARE Classic, NDEF, Samsung S25'

def p(text): return doc.add_paragraph(text)
def h(text): doc.add_heading(text,2)
def step(n,text): p(str(n)+'. '+text)
def page(title): doc.add_page_break(); doc.add_heading(title,1)
def code(text):
    para=p(text); para.paragraph_format.space_after=Pt(8)
    for r in para.runs: r.font.name='Consolas';r.font.size=Pt(9)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ridx,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); nr=OxmlElement('w:cantSplit');trpr.append(nr)
        if ridx==0: repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for i,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[i])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'16324F' if ridx==0 else ('F0F4F7' if ridx%2==0 else 'FFFFFF'));pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'95');x.set(qn('w:type'),'dxa');margins.append(x)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                for r in para.runs:r.font.size=Pt(9);r.bold=ridx==0;r.font.color.rgb=RGBColor(255,255,255) if ridx==0 else RGBColor(0,0,0)
    p('')
    return t

# 1
p('NFC Card Lab').style='Title'
p('Dokumentacja projektu i instrukcja obsługi').style='Subtitle'
p('Stan projektu na 5 października 2026 • Aplikacja 1.0.0 • Dokumentacja 1.1')
p('Aktualna wersja ma prosty polski interfejs Tag, Zapis i Przygotuj oraz wybór aplikacji z danego telefonu. Każdy użytkownik wybiera własną aplikację i zapisuje jej nazwę pakietu na swoim tagu.')
p('NFC Card Lab to lokalna aplikacja Android do diagnostyki własnych tagów NFC, pracy z pamięcią MIFARE Classic oraz odczytu, zapisu i formatowania NDEF. Korzysta z NFC telefonu; nie wymaga zewnętrznego czytnika, serwera ani chmury.')
p('Na Samsungu Galaxy S25 potwierdzono konwersję własnego taga o UID 42 27 04 26 do NDEF oraz zapis rekordu URI https://bitback.pl/ z porównaniem odczytanych bajtów. Wynik dotyczy tego telefonu i tego taga. Samo uruchomienie przeglądarki po przyłożeniu taga pozostaje do sprawdzenia.')
table(['Parametr','Wartość'],[
('Urządzenie docelowe','Samsung Galaxy S25 SM-S931B • Android 16'),
('Pakiet aplikacji','pl.misieklab.nfccardlab'),('Wersja','1.0.0 • versionCode 1 • APK debug'),
('SDK','minSdk 29 • compileSdk 36 • targetSdk 36'),
('Technologie','Kotlin • Jetpack Compose • Material 3 • coroutines • ViewModel • StateFlow'),
('Komunikacja','Android NFC API • Reader Mode • bez Internet permission'),
('Projekt lokalny',r'D:\Codex\NFC-Card-Lab')],[4.3,12.7])
h('Zakres dokumentacji')
p('Rozdziały obejmują obsługę aplikacji, bezpieczny zapis bloków, NDEF, zaawansowane operacje na tagach Bambu, architekturę, budowanie i instalację, wyniki testów oraz lokalizacje dowodów. Dokument powstał z aktualnych źródeł, raportów testów i zapisanej walidacji urządzenia.')
p('Zasada nadrzędna: Block 0 pozostaje zablokowany dla bezpośrednich zapisów. Zmiana trailerów jest operacją zaawansowaną i wymaga osobnego zezwolenia. Aplikacja nie implementuje klonowania kart dostępu, brute force, ataków na klucze ani zmiany UID.')

# 2
page('1 Uruchomienie i podstawowa obsługa')
step(1,'Włącz NFC w ustawieniach telefonu i odblokuj ekran. Uruchom NFC Card Lab.')
step(2,'W zakładce Tag przyłóż własny tag do tylnej części telefonu. Poczekaj na typ i identyfikator. Podczas pracy trzymaj tag nieruchomo.')
step(3,'Jeżeli tag ma NDEF, przejdź do Zapis. Wybierz Link, Tekst albo Aplikacja. Zapisz i sprawdź pokaże potwierdzenie, zapisze i porówna ponowny odczyt.')
step(4,'Jeżeli tag wymaga przygotowania, użyj Przygotuj. Utrata kontaktu wymaga ponownego przyłożenia. Niepotwierdzony zapis trzeba sprawdzić przed powtórzeniem.')
table(['Zakładka','Zastosowanie'],[
('Tag','Typ, identyfikator, gotowość do zapisu i odczytana treść.'),
('Zapis','Link, tekst lub wybrana aplikacja; jeden przycisk Zapisz i sprawdź.'),
('Przygotuj','Formatowanie pustego taga lub prowadzona obsługa własnego taga Bambu.')],[3.1,13.9])
p('Ustawienia → Tryb zaawansowany udostępniają SCAN, DETAILS, MEMORY, WRITE, NDEF, APDU i LOG, klucze, HEX oraz eksport. Wybór interfejsu jest zapamiętywany; uprawnienia zapisu trailerów pozostają osobne.')
h('Dodaj aplikację z własnego telefonu')
step(1,'W Zapis wybierz Aplikacja, a następnie Dodaj aplikację. Wyszukaj i wybierz aplikację z lokalnej listy. Możesz ją wybrać jeszcze przed przyłożeniem taga.')
step(2,'Przyłóż zapisywalny tag NDEF i wybierz Zapisz i sprawdź. Potwierdź zastąpienie obecnej wiadomości. Na tag trafi nazwa pakietu, nie instalator aplikacji.')
step(3,'Po zakończeniu wyjdź z NFC Card Lab, odsuń tag i ponownie przyłóż go do odblokowanego telefonu. Próba otwarcia aplikacji nadal wymaga testu z tagiem.')
h('Otwieranie zapisanego linku')
p('Na tagu 42 27 04 26 zapisano https://bitback.pl/. Aby sprawdzić uruchomienie strony, wyjdź z NFC Card Lab na ekran główny, pozostaw NFC włączone, odblokuj telefon, odsuń tag i przyłóż go ponownie. Aktywny Reader Mode aplikacji służy diagnostyce i przejmuje skanowanie.')
p('Android 16 kieruje linki HTTP/HTTPS z NDEF przez ACTION_VIEW. Chrome powinien obsłużyć link, jeśli jest domyślną przeglądarką; wybór aplikacji zależy od konfiguracji telefonu. Strona WWW wymaga dostępu do sieci w przeglądarce, mimo że NFC Card Lab działa bez Internetu. Źródło: Android Developers [1].')

# 3
page('2 Pamięć i bezpieczny zapis MIFARE Classic')
h('Mapa pamięci i uwierzytelnianie')
p('Dostęp do pamięci jest oferowany, gdy MifareClassic.get(tag) zwraca technologię. Rozmiar, typ oraz układ sektorów pochodzą z API. W Classic 1K występuje 16 sektorów po 4 bloki; każdy blok ma 16 bajtów. Ostatni blok sektora jest trailerem. Większe tagi mogą mieć inny układ sektorów.')
step(1,'W MEMORY lub WRITE wybierz sektor i Key A albo Key B. Wprowadź posiadany klucz, dokładnie 6 bajtów: FF FF FF FF FF FF lub FFFFFFFFFFFF. Oba formaty są poprawne.')
step(2,'AUTHENTICATE potwierdza wyłącznie podany klucz w wybranym sektorze. TRY DEFAULT KEY próbuje tylko MifareClassic.KEY_DEFAULT. Brak klucza lub błędny klucz nie jest powodem do uruchamiania ataku.')
step(3,'READ ALL ACCESSIBLE przechodzi po sektorach z jednym wybranym kluczem. AUTH FAILED w sektorze powoduje przejście dalej. Utrata taga może zakończyć cały odczyt jako niekompletny.')
h('Zapis bloku danych')
step(1,'Wybierz sektor i bezwzględny numer bloku należący do tego sektora, np. sektor 1 i blok 4.')
step(2,'Wpisz dokładnie 16 bajtów HEX. Edytor obsługuje duże litery, odstępy, licznik bajtów, wklejanie i czyszczenie. Niepełne lub zbyt długie dane blokują zapis.')
code('01 02 03 04 05 06 07 08\n09 0A 0B 0C 0D 0E 0F 10')
step(3,'READ BLOCK pozwala sprawdzić zawartość. WRITE / WRITE + VERIFY przygotowują potwierdzenie z UID, bieżącymi i nowymi bajtami. Zapis w tej wersji zawsze obejmuje weryfikację.')
step(4,'Potwierdź WRITE. Aplikacja łączy się, uwierzytelnia, sprawdza, czy blok nie zmienił się od potwierdzenia, zapisuje, odczytuje i porównuje bajty. Wynik: WRITE VERIFIED albo VERIFY FAILED.')
h('Ochrona struktury taga')
p('Block 0 jest zawsze chroniony przed bezpośrednim zapisem. Trailer jest wyznaczany przez sectorToBlock() i getBlockCountInSector(), nie przez sztywną listę. Domyślnie jego zapis jest blokowany. Nawet poprawny blok danych może zawierać istotną część NDEF lub innej aplikacji; wybieraj wyłącznie pamięć przeznaczoną do zmiany.')
p('Trailer dzieli się na Key A (bajty 0–5), access bits (6–8), GPB (9) i Key B (10–15). Odczytane zera w polu klucza nie dowodzą, że klucz jest zerowy. Aplikacja maskuje pola kluczy i nie przedstawia ich jako wiarygodnego odczytu.')

# 4
page('3 Odczyt zapis i formatowanie NDEF')
p('NDEF jest standardowym formatem rekordów na tagach NFC. Zwykły zapis wykonuje Ndef.writeNdefMessage(); aplikacja nie zapisuje adresu WWW jako przypadkowego ciągu ASCII w dowolnym bloku.')
table(['Rekord','Dane użytkownika','Przykład'],[
('TEXT','Treść UTF-8 i kod języka','Test NFC • język pl'),
('URI','Pełny adres URI bez białych znaków','https://bitback.pl/'),
('MIME','Typ MIME i treść UTF-8','text/plain • Test NFC'),
('EXTERNAL','domain:type i treść UTF-8','example.com:lab • Test NFC'),
('APPLICATION','Aplikacja wybrana z telefonu','Android Application Record z pakietem')],[2.5,7.4,7.1])
h('Zapis na tagu NDEF')
step(1,'W prostym widoku użyj Zapis i wybierz treść. W trybie zaawansowanym otwórz NDEF, sprawdź Writable i Capacity oraz wybierz rodzaj rekordu.')
step(2,'Aplikacja tworzy NdefRecord oraz NdefMessage i pokazuje rozmiar wiadomości. isWritable musi być prawdziwe, a wiadomość nie może przekraczać maxSize.')
step(3,'Wybierz Zapisz i sprawdź lub zaawansowany WRITE NDEF + VERIFY. Sprawdź UID i treść w dialogu. Zapis zastępuje istniejącą wiadomość NDEF pojedynczą nowo utworzoną wiadomością.')
step(4,'Po potwierdzeniu aplikacja odczytuje Ndef.ndefMessage i porównuje serializowane bajty z zapisanym komunikatem. NDEF WRITE VERIFIED oznacza zgodny odczyt w tej operacji.')
h('Formatowanie zgodnego pustego taga')
p('Jeżeli tag nie udostępnia Ndef, lecz udostępnia NdefFormatable, pojawia się TAG IS NOT NDEF FORMATTED i FORMAT AS NDEF. Po osobnym potwierdzeniu aplikacja używa connect() oraz format() z minimalnym poprawnym rekordem Text.')
p('Zakończenie komendy formatowania nie oznacza jeszcze weryfikacji. Po zamknięciu połączenia trzeba odsunąć tag i przyłożyć go ponownie. Android musi wykryć technologię Ndef. Nie każdy tag oznaczony NdefFormatable da się sformatować przez dany telefon.')
p('Standardowe NdefFormatable.format() nie zostało skutecznie potwierdzone na testowanych tagach Bambu. Sukces opisany w tej dokumentacji dotyczy osobnej, jawnie zatwierdzonej konwersji Classic do NDEF, opisanej w następnym rozdziale.')
h('Błędy operacji')
p('TAG LOST wymaga ponownego przyłożenia taga. READ ONLY i MESSAGE TOO LARGE blokują zapis przed I/O. FORMAT NOT VERIFIED oznacza brak potwierdzenia po ponownym skanie. NOT SUPPORTED BY ANDROID NFC API oznacza niedostępną technologię lub funkcję, a nie atrapę operacji.')

# 5
page('4 Zaawansowane operacje na własnym tagu Bambu')
p('Operacje te są przeznaczone dla własnego taga Bambu zgodnego z MIFARE Classic 1K, z UID 4 B. Używają opublikowanego deterministycznego wyprowadzenia kluczy A/B z UID [4]. Nie wyszukują kluczy brute force. Przed zapisem wszystkie 16 sektorów musi przejść uwierzytelnienie i pełny odczyt.')
h('Warunki uruchomienia')
p('W prostym widoku Przygotuj osobno potwierdź własność taga Bambu, utratę danych oraz zmianę ustawień sektorów. Po sprawdzonym czyszczeniu konwersja jest dostępna wyłącznie dla tego samego UID. Zgoda dotyczy danej operacji; nie odblokowuje ogólnego edytora trailerów. NDEF pomija ten proces. Typ Classic 1K sam nie dowodzi pochodzenia Bambu.')
p('W trybie zaawansowanym w SETTINGS włącz Developer mode oraz Allow sector trailer writes. Operacja wymaga dodatkowego dialogu i wpisania ERASE lub CONVERT oraz UID bez odstępów. Zgoda jest przypisana do bieżącej sesji taga. Opuszczenie aplikacji unieważnia sesję i wyłącza ustawienia deweloperskie.')
h('Kontrolowane czyszczenie danych')
p('ERASE BAMBU DATA w MEMORY zapisuje kopię wszystkich 64 bloków, z maskowaniem kluczy. Kopia musi być utrwalona na dysku przed pierwszą zmianą taga. Następnie operacja tymczasowo zmienia wymagane uprawnienia danych, zeruje 47 bloków danych i przywraca pierwotne uprawnienia w finally.')
p('Odczyt końcowy sprawdza wszystkie 64 bloki: dane zerowe, niezmieniony Block 0, kontrolne bajty trailerów i uwierzytelnienie obu pierwotnych kluczy. Samo wyzerowanie danych nie formatuje taga jako NDEF.')
h('Konwersja do NDEF')
p('CONVERT CLEARED TAG TO NDEF w NDEF wymaga uprzednio wyczyszczonych 47 bloków. Niezerowe dane powodują odmowę przed zapisem. Aplikacja tworzy MAD1, CRC, mapowanie sektorów NFC i poprawny NDEF TLV zgodnie z NXP [2, 3]. Klucze A zmieniają się na standardowe publiczne klucze MAD/NFC; pierwotne klucze B i Block 0 pozostają zachowane.')
p('NFC sektory danych są przygotowywane i sprawdzane przed publikacją MAD w sektorze 0. Każdy zapis jest odczytywany. Końcowo sprawdzane są 64 bloki, dostęp nowymi kluczami A i dotychczasowymi B. Po sukcesie wymagane jest ponowne przyłożenie oraz wykrycie Ndef przez Android.')
table(['Sektor','Publiczny Key A','Access bits i GPB'],[
('0 • MAD','A0 A1 A2 A3 A4 A5','78 77 88 C1'),
('1–15 • NFC','D3 F7 D3 F7 D3 F7','7F 07 88 40')],[3.4,7.2,6.4])
p('W razie błędu konwerter próbuje odtworzyć poprzedni stan za pomocą niezmienionych kluczy B. TAG LOST lub zamknięcie połączenia może uniemożliwić rollback. ROLLBACK INCOMPLETE oznacza stan częściowy: zachowaj kopię i dziennik; nie ponawiaj konwersji w ciemno. Pełne odzyskiwanie po utracie taga nie jest automatyczne.')

# 6
page('5 Architektura i bezpieczeństwo aplikacji')
p('SimpleScreen.kt obsługuje widok podstawowy, a SimpleWorkflow.kt komunikaty i warunki przejścia do konwersji. AppPicker.kt wczytuje i filtruje lokalne aplikacje poza wątkiem głównym.')
p('MainActivity obsługuje cykl życia Reader Mode. LabViewModel publikuje LabState przez StateFlow i koordynuje zadania. Composable wyświetlają stan i przekazują intencje użytkownika; transport i operacje taga pozostają w osobnej warstwie nfc.')
table(['Moduł','Odpowiedzialność'],[
('NfcManager.kt','Reader Mode, token sesji, Mutex, Dispatchers.IO, connect/close, unieważnianie starego taga.'),
('NfcTechDetector.kt','Dynamiczna klasyfikacja i parametry publicznego Android API.'),
('NfcTagInfo.kt','Modele informacji taga, sektorów, bloków i NDEF.'),
('NfcHexUtils.kt','Parser/formatter HEX i UID, walidacja długości oraz BlockPolicy i układ sektorów.'),
('NfcResult.kt','Wynik operacji i mapowanie błędów NFC na komunikaty.'),
('MifareClassicManager.kt','Autoryzacja A/B, blokowe I/O, dump, porównanie i polityka zapisu; ClassicPort do testów.'),
('NdefManager.kt','Tworzenie/odczyt/zapis rekordów, pojemność i standardowe formatowanie.'),
('ControlledErase.kt','Deterministyczne klucze Bambu, preflight, zerowanie i odtworzenie uprawnień.'),
('ClassicNdefConversion.kt','MAD1, CRC, TLV, konwersja z weryfikacją i próbą rollbacku.'),
('DumpExporter.kt','Eksport TXT/JSON z maskowaniem pól kluczy w trailerach.')],[5.2,11.8])
h('Cykl życia i serializacja')
p('Reader Mode skanuje NFC-A/B/F/V i nie pomija platformowego sprawdzania NDEF. I/O działa poza wątkiem głównym. Wspólny Mutex i bramka zajętości zapobiegają równoległym operacjom. Token generacji wiąże potwierdzenie zapisu z aktualnym tagiem. onPause anuluje operacje, unieważnia sesję i zamyka technologię; otwarcia są zamykane w finally także po błędzie connect().')
h('Klucze uprawnienia i dane lokalne')
p('Manifest udostępnia wyszukiwanie MAIN/LAUNCHER przez queries. Nie dodano QUERY_ALL_PACKAGES ani nowych uprawnień. Lista aplikacji nie jest eksportowana ani logowana.')
p('Klucze są przechowywane w pamięci, nigdy w logach, analytics ani eksporcie. Tymczasowe tablice bajtów są zerowane; niemodyfikowalne String JVM nie gwarantują natychmiastowego wymazania. Brak serwera, SDK raportowania awarii oraz Internet permission. Manifest wymaga NFC; backup aplikacji jest wyłączony. FLAG_SECURE ogranicza zrzuty ekranów i podgląd w ostatnich aplikacjach.')
p('Obsługa błędów obejmuje IOException, TagLostException, FormatException, IllegalArgumentException i SecurityException. Log jest lokalnym ograniczonym buforem zdarzeń. Kopie operacji zaawansowanych i dzienniki faz są plikami prywatnymi aplikacji.')

# 7
page('6 Budowanie instalacja i obsługa problemów')
h('Konfiguracja projektu')
table(['Składnik','Wersja używana w projekcie'],[
('Android SDK','API 36 stabilne • Build Tools 36.0.0 • minimum API 29'),
('Gradle i Android plugin','Gradle 8.12 • Android Gradle Plugin 8.10.1'),
('Kotlin i JVM','Kotlin 2.0.21 • docelowa JVM 17'),
('Compose i Material','Compose BOM 2024.12.01 • Material 3 z BOM'),
('Lifecycle i coroutines','Lifecycle 2.8.7 • coroutines 1.9.0'),
('Testy','JUnit 4.13.2 • Robolectric 4.16.1')],[5,12])
p('Otwórz katalog projektu w Android Studio. local.properties powinien wskazywać zainstalowane Android SDK. Projekt korzysta z własnego Gradle Wrapper; nie wymaga systemowego Gradle. Wersje są przypięte i opisują faktyczny stan projektu, nie deklarację najnowszych bibliotek.')
code('$env:JAVA_HOME = "C:\\Program Files\\Android\\Android Studio\\jbr"\n.\\gradlew.bat test\n.\\gradlew.bat assembleDebug\n.\\gradlew.bat lintDebug')
p('APK: app/build/outputs/apk/debug/app-debug.apk. Jest podpisany kluczem debug. Instalacja aktualizacji wymaga zgodnego podpisu; nie usuwaj poprzedniej aplikacji w ciemno w razie INSTALL_FAILED_UPDATE_INCOMPATIBLE, bo usunięcie może skasować jej prywatne kopie.')
h('Instalacja przez ADB')
code('adb devices -l\nadb -s <adres:port> install -r --user 0 app-debug.apk\nadb -s <adres:port> shell am start -n pl.misieklab.nfccardlab/.MainActivity')
p('Debugowanie bezprzewodowe wymaga sparowania komputera i telefonu. Port połączenia może się zmieniać; używaj aktualnego adresu z ustawień telefonu lub adb mdns services. Nie zapisuj kodu parowania w dokumentacji. Aplikacja nie wymaga ADB po instalacji.')
h('Typowe problemy')
p('Brak taga: odblokuj ekran, włącz NFC, popraw położenie taga. AUTH FAILED: sprawdź sektor, rodzaj i dokładne 6 bajtów posiadanego klucza; KEY_DEFAULT nie pasuje do każdego taga. PROTECTED BLOCK: wybierz blok danych; Block 0 nie można odblokować. Po zmianie aplikacji lub TAG LOST ponów skan i potwierdzenie. VERIFY FAILED: odczytaj faktyczną zawartość przed kolejną zmianą.')

# 8
page('7 Wyniki testów i zakres potwierdzenia')
p('Stan obejmuje kolejne etapy z 5 października 2026. Testy fizycznego czyszczenia i URI dotyczą wcześniejszego APK; interfejs i lista aplikacji były sprawdzone po aktualizacji. Tworzenie tej dokumentacji nie wykonuje nowych operacji NFC ani nowych zapisów na tagu.')
table(['Funkcja','Potwierdzenie i ograniczenie'],[
('BUILD_STATUS','BUILD SUCCESSFUL • test assembleDebug lintDebug'),
('TEST_STATUS','91 testów debug + 91 release • 0 błędów i pominięć'),
('NFC_READER_MODE','DEVICE_TESTED • S25 • UID 12 66 FA 25 i 42 27 04 26'),
('NDEF_READ i NDEF_WRITE','DEVICE_TESTED • URI https://bitback.pl/ • NDEF WRITE VERIFIED'),
('NDEF_FORMAT','Standard format() niepotwierdzony; custom Classic conversion DEVICE_TESTED na UID 42 27 04 26'),
('MIFARE_CLASSIC_READ','DEVICE_TESTED • 64 bloki odczytane podczas operacji kontrolowanych'),
('MIFARE_CLASSIC_AUTH','DEVICE_TESTED • A+B wszystkie 16 sektorów; po konwersji publiczny A i pierwotny B'),
('MIFARE_CLASSIC_WRITE','DEVICE_TESTED • kontrolowane zerowanie i konwersja konkretnego taga'),
('WRITE_VERIFY','DEVICE_TESTED • 47 zerowych bloków, struktura konwersji oraz odczyt NDEF'),
('PROTECTED_BLOCKS','UNIT_TESTED • bezpośredni Block 0 zawsze zablokowany; trailery domyślnie'),
('MEMORY_DUMP','UNIT/COMPILE_TESTED; ogólny przycisk READ ALL ACCESSIBLE wymaga osobnego testu UI'),
('EXPORT_JSON i TXT','UNIT_TESTED • maskowanie kluczy; picker dokumentów REQUIRES_DEVICE_TEST'),
('Interfejs i instalacja','DEVICE_TESTED • aktualizacja S25, NFC aktywne, lista aplikacji, wyszukiwanie i wybór'),
('APPLICATION zapis i start','UNIT/COMPILE_TESTED • rekord AAR; zapis na tagu i start aplikacji REQUIRES_DEVICE_TEST'),
('Otwarcie Chrome','REQUIRES_DEVICE_TEST • zapis URI potwierdzony, start przeglądarki nie')],[5.3,11.7])
h('Znaczenie oznaczeń')
p('COMPILE_TESTED oznacza udaną kompilację. UNIT_TESTED oznacza testy JVM/Robolectric i transportów zastępczych. DEVICE_TESTED oznacza faktyczną operację na telefonie i tagu, w opisanym zakresie. REQUIRES_DEVICE_TEST oznacza, że implementacja lub opis nie jest dowodem działania sprzętu.')
p('Testy obejmują parser/formatter HEX i UID, klucze 6 B i dane 16 B, sektory/trailery, kontrolę uprawnień i porównania zapisu, NDEF Text/URI/MIME/External/Application i pojemność, redakcję eksportu, sesje/close oraz cykl Activity. Dodatkowo badano HKDF, preflight przed zmianami, błędy zapisu, MAD CRC i rollback konwersji.')

# 9
page('8 Artefakty dowody i utrzymanie projektu')
h('Aktualny APK i projekt')
code(r'Projekt: D:\Codex\NFC-Card-Lab'+'\n'+r'APK: app\build\outputs\apk\debug\app-debug.apk')
p('SHA-256 aktualnego APK z opcją Dodaj aplikację:')
apk_sha=hashlib.sha256((root/'app/build/outputs/apk/debug/app-debug.apk').read_bytes()).hexdigest().upper()
code(apk_sha[:32]+'\n'+apk_sha[32:])
p('Hash jest podzielony na dwa wiersze wyłącznie dla czytelności; należy go połączyć bez spacji. Przy dalszych zmianach buduj nowy APK i zapisuj nowy hash. Pliki README.md, VALIDATION.md oraz DEVICE_TEST_CHECKLIST.md uzupełniają tę dokumentację. Wyniki 61, 75 i 83 testów w starszych wpisach są historyczne. Aktualny etap ma 91 testów dla każdego wariantu.')
h('Dowody pracy z tagiem')
code('backups/tag-42270426-erase-20261005/\nbackups/tag-42270426-ndef-20261005/')
p('W drugim folderze znajdują się convert-42270426-1791218008542-before.json, -after.json i -state.txt oraz ndef-uri-device-evidence.json. Kopie przed/po i dziennik skopiowano binarnie z telefonu; SHA-256 każdego pliku zgadzał się z oryginałem. Dane kluczy są zamaskowane. after.json opisuje pusty rekord Text po konwersji, przed zapisem URI; nie jest końcowym dumpem linku.')
p('Zapisany dowód UI z 5 października 2026 godz. 18:38 czasu Europe/Warsaw pokazuje NDEF WRITE VERIFIED, URI https://bitback.pl/, writable YES i 716 B. Wiadomość URI zajmuje 16 B. Kopie na telefonie są w files/erase-backups i files/conversion-backups wewnątrz prywatnego katalogu aplikacji.')
h('Ograniczenia i dalsza walidacja')
p('Brak gwarancji zgodności wszystkich tagów i telefonów z MIFARE Classic. Przerwany zapis może pozostawić stan częściowy. Nie ma automatycznej naprawy po utracie taga, dowolnego APDU, odzyskiwania kluczy atakami ani zapamiętywania kluczy. MIME/External przyjmują tekst UTF-8, nie pliki binarne. Aplikacja nie jest wydaniem sklepowym. Sam zapis linku nie potwierdza działania serwisu WWW.')
p('Przed kolejnym wydaniem sprawdź zapis APPLICATION i otwarcie wybranej aplikacji po ponownym przyłożeniu, Chrome, eksport, Text i standardowe formatowanie na zgodnym pustym tagu. Powtórz testy i lint po zmianach kodu. Projekt nie jest w zatwierdzonym manifeście synchronizacji; nie utworzono ani nie opublikowano repozytorium Git. Lokalna dokumentacja nie zmienia tego statusu.')
h('Źródła techniczne')
for text in [
'[1] Android Developers NFC basics • developer.android.com/develop/connectivity/nfc/nfc',
'[2] NXP AN1305 • nxp.com/docs/en/application-note/AN1305.pdf',
'[3] NXP AN10787 • nxp.com/docs/en/application-note/AN10787.pdf',
'[4] Publiczne wyprowadzenie kluczy Bambu • github.com/queengooborg/Bambu-Lab-RFID-Tag-Guide/blob/main/deriveKeys.py',
'[5] Android Package visibility • developer.android.com/training/package-visibility/declaring']:
    para=p(text)
    for run in para.runs:run.font.size=Pt(9)

# 10
page('9 Schemat działania i historia zmian')
h('Od rozpoznania do zapisu')
step(1,'Rozpoznanie: przyłóż tag. Android podaje technologie, UID oraz możliwości pamięci. Aplikacja nie zgaduje producenta układu z samego SAK.')
step(2,'Tag ma NDEF: przejdź do Zapis. Czyszczenie i konwersja Bambu nie są potrzebne. Nowa wiadomość zastępuje starą.')
step(3,'Własny tag Bambu bez NDEF: potwierdź czyszczenie i zmianę ustawień. Preflight sprawdza klucze i pamięć; lokalny odczyt jest zapisany przed pierwszą zmianą. Czyszczenie zeruje 47 bloków danych i sprawdza wynik.')
step(4,'Po sprawdzonym czyszczeniu tego samego UID potwierdź Przygotuj do linków. Konwersja tworzy strukturę MAD/NDEF i sprawdza bloki. Odsuń i przyłóż tag ponownie: dopiero wykrycie Ndef potwierdza format.')
step(5,'Inny pusty tag z NdefFormatable: osobno potwierdź standardowe formatowanie Androida. Powodzenie zależy od taga i telefonu. Brak obsługiwanej technologii oznacza brak dostępnej metody.')
step(6,'Zapis: wybierz Link, Tekst lub Aplikacja. Pojemność i zapisowalność są sprawdzane przed zapisem. Potwierdzenie dotyczy bieżącego taga; wynik wymaga zgodnego ponownego odczytu.')
h('Wybór aplikacji na różnych telefonach')
p('Każdy użytkownik wybiera aplikację ze swojego telefonu i zapisuje ją na swoim tagu. Nie trzeba znać jego listy aplikacji. Jeden tag wskazuje jeden zapisany pakiet; nie dobiera innej aplikacji automatycznie dla każdego telefonu.')
p('Lista obejmuje aktywne aplikacje z ekranem uruchamiania w bieżącym profilu. Ukryte, zawieszone lub należące do innego profilu mogą nie być widoczne. Android obsługuje AAR [1]. Na telefonie bez wybranego pakietu może otworzyć Google Play. Aplikacja lub ustawienia systemu mogą wymagać dodatkowej zgody.')
h('Kolejne etapy tego samego dnia')
table(['Etap','Testy na wariant','Potwierdzenie urządzenia'],[
('Konwersja i URI','75','Classic 1K UID 42 27 04 26, NDEF 716 B, URI 16 B'),
('Prosty interfejs','83','Instalacja, trzy zakładki i aktywny czytnik po ponownym połączeniu ADB'),
('Dodaj aplikację','91','Aktualizacja, lokalna lista, wyszukiwanie i wybór; bez zapisu AAR na tagu')],[4.4,3.4,9.2])
p('Starsze źródła i APK zachowano w backups/before-simple-ui-20261005/ oraz backups/before-app-record-20261005/. Poprzednia dokumentacja znajduje się w backups/before-docs-update-20261005/. Wersja dokumentacji 1.1 opisuje aktualny stan bez przypisywania nowych funkcji wcześniejszym testom sprzętu.')

# Remove inherited Word template borders and implicit outline numbering.
for style in doc.styles:
    for element in list(style.element.iter(qn('w:pBdr'))): element.getparent().remove(element)
for para in doc.paragraphs:
    for element in list(para._p.iter(qn('w:pBdr'))): element.getparent().remove(element)
    if para.style.name.startswith('Heading'):
        pr=para._p.get_or_add_pPr()
        for old in list(pr.findall(qn('w:numPr'))): pr.remove(old)
        np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);pr.append(np)
        para.paragraph_format.left_indent=Cm(0)
        para.paragraph_format.first_line_indent=Cm(0)
doc.save(out)
print(out)
print('SHA256='+hashlib.sha256(out.read_bytes()).hexdigest())
